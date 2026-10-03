#!/usr/bin/env python3
"""Extract plain text from .txt/.md/.docx/.pdf (CVs, drafts, LinkedIn exports).

Usage: python read_doc.py <file> [--out text.txt]
.docx needs `pip install python-docx`; .pdf needs `pip install pypdf`.
Both are optional, the script explains what to install if missing.
"""
import argparse
import sys


def read_document(path):
    low = path.lower()
    if low.endswith(".docx"):
        try:
            import docx  # python-docx
        except ImportError:
            raise SystemExit("Reading .docx needs python-docx:  pip install python-docx  (or paste the text / save as .txt)")
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(p for p in parts if p.strip())
    if low.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError:
            raise SystemExit("Reading .pdf needs pypdf:  pip install pypdf  (or paste the text / save as .txt)")
        reader = PdfReader(path)
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    with open(path, encoding="utf-8") as f:
        return f.read()


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--out")
    a = ap.parse_args()
    text = read_document(a.path)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"wrote {len(text)} chars to {a.out}")
    else:
        print(text)


if __name__ == "__main__":
    main()
